"""Reporting dashboard."""
from dashboard.views._common import *


# =============================================================================
# Reporting Dashboard
# =============================================================================

@login_required
def reports_view(request):
    from reference.models import Operator
    operators = Operator.objects.filter(enabled=True).order_by('name')
    return render(request, 'dashboard/reports.html', {'operators': operators})


@login_required
def reports_api(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'POST required'}, status=405)

    from django.db.models.functions import TruncDate
    from django.db.models import ExpressionWrapper, DurationField
    from collection.models import DistributionLog, ProcessingError
    from portals.models import OutputPortal, DistributionRule

    start_date = request.POST.get('start_date', '').strip()
    end_date = request.POST.get('end_date', '').strip()
    operator = request.POST.get('operator', '').strip()

    if not start_date or not end_date:
        return JsonResponse({'success': False, 'message': 'Both start_date and end_date are required.'}, status=400)

    try:
        dt_start = datetime.strptime(start_date, '%Y-%m-%d')
        dt_end = datetime.strptime(end_date, '%Y-%m-%d') + timedelta(days=1)
    except ValueError:
        return JsonResponse({'success': False, 'message': 'Invalid date format. Use YYYY-MM-DD.'}, status=400)

    period_days = (dt_end - dt_start).days
    prev_end = dt_start
    prev_start = prev_end - timedelta(days=period_days)

    file_qs = CDRFile.objects.filter(created_at__gte=dt_start, created_at__lt=dt_end)
    prev_file_qs = CDRFile.objects.filter(created_at__gte=prev_start, created_at__lt=prev_end)

    if operator:
        file_qs = file_qs.filter(operator_code=operator)
        prev_file_qs = prev_file_qs.filter(operator_code=operator)

    # --- Input: CDRFiles grouped by DataSource ---
    def _build_input_sources(qs):
        input_sources = []
        sources = DataSource.objects.filter(enabled=True).order_by('name')
        for src in sources:
            src_files = qs.filter(source=src)
            agg = src_files.aggregate(
                total_files=Count('id'),
                total_records=Sum('records_valid'),
                completed=Count('id', filter=Q(status=CDRFile.Status.COMPLETED)),
                failed=Count('id', filter=Q(status=CDRFile.Status.FAILED)),
                duplicates=Count('id', filter=Q(status=CDRFile.Status.DUPLICATE)),
                first_received=Min('created_at'),
                last_received=Max('created_at'),
            )
            if agg['total_files'] == 0:
                continue
            by_stream = list(
                src_files.values('decoder_type')
                .annotate(files=Count('id'), records=Sum('records_valid'))
                .order_by('decoder_type')
            )
            input_sources.append({
                'source_id': src.pk,
                'source_name': src.name,
                'total_files': agg['total_files'] or 0,
                'total_records': agg['total_records'] or 0,
                'completed': agg['completed'] or 0,
                'failed': agg['failed'] or 0,
                'duplicates': agg['duplicates'] or 0,
                'first_received': agg['first_received'].strftime('%Y-%m-%d %H:%M') if agg['first_received'] else '',
                'last_received': agg['last_received'].strftime('%Y-%m-%d %H:%M') if agg['last_received'] else '',
                'by_stream': [{'stream': r['decoder_type'] or '-',
                               'files': r['files'], 'records': r['records'] or 0}
                              for r in by_stream],
            })

        no_src_files = qs.filter(source__isnull=True)
        no_src_agg = no_src_files.aggregate(
            total_files=Count('id'),
            total_records=Sum('records_valid'),
            completed=Count('id', filter=Q(status=CDRFile.Status.COMPLETED)),
            failed=Count('id', filter=Q(status=CDRFile.Status.FAILED)),
            duplicates=Count('id', filter=Q(status=CDRFile.Status.DUPLICATE)),
            first_received=Min('created_at'),
            last_received=Max('created_at'),
        )
        if no_src_agg['total_files']:
            no_src_streams = list(
                no_src_files.values('decoder_type')
                .annotate(files=Count('id'), records=Sum('records_valid'))
                .order_by('decoder_type')
            )
            input_sources.append({
                'source_id': None,
                'source_name': '(No Source)',
                'total_files': no_src_agg['total_files'] or 0,
                'total_records': no_src_agg['total_records'] or 0,
                'completed': no_src_agg['completed'] or 0,
                'failed': no_src_agg['failed'] or 0,
                'duplicates': no_src_agg['duplicates'] or 0,
                'first_received': no_src_agg['first_received'].strftime('%Y-%m-%d %H:%M') if no_src_agg['first_received'] else '',
                'last_received': no_src_agg['last_received'].strftime('%Y-%m-%d %H:%M') if no_src_agg['last_received'] else '',
                'by_stream': [{'stream': r['decoder_type'] or '-',
                               'files': r['files'], 'records': r['records'] or 0}
                              for r in no_src_streams],
            })
        return input_sources

    input_sources = _build_input_sources(file_qs)

    # --- Output: DistributionLogs grouped by OutputPortal ---
    dist_qs = DistributionLog.objects.filter(delivered_at__gte=dt_start, delivered_at__lt=dt_end)
    prev_dist_qs = DistributionLog.objects.filter(delivered_at__gte=prev_start, delivered_at__lt=prev_end)
    if operator:
        dist_qs = dist_qs.filter(cdr_file__operator_code=operator)
        prev_dist_qs = prev_dist_qs.filter(cdr_file__operator_code=operator)

    output_portals = []
    for portal in OutputPortal.objects.filter(is_active=True).order_by('name'):
        p_logs = dist_qs.filter(output_portal=portal)
        agg = p_logs.aggregate(
            total_files=Count('id'),
            total_records=Sum('record_count'),
            success=Count('id', filter=Q(status=DistributionLog.Status.SUCCESS)),
            failed=Count('id', filter=Q(status=DistributionLog.Status.FAILED)),
            skipped=Count('id', filter=Q(status=DistributionLog.Status.SKIPPED)),
            first_delivered=Min('delivered_at'),
            last_delivered=Max('delivered_at'),
        )
        if agg['total_files'] == 0:
            continue
        by_stream = list(
            p_logs.filter(cdr_file__isnull=False)
            .values('cdr_file__decoder_type')
            .annotate(files=Count('id'), records=Sum('record_count'))
            .order_by('cdr_file__decoder_type')
        )
        output_portals.append({
            'portal_id': portal.pk,
            'portal_name': portal.name,
            'total_files': agg['total_files'] or 0,
            'total_records': agg['total_records'] or 0,
            'success': agg['success'] or 0,
            'failed': agg['failed'] or 0,
            'skipped': agg['skipped'] or 0,
            'first_delivered': agg['first_delivered'].strftime('%Y-%m-%d %H:%M') if agg['first_delivered'] else '',
            'last_delivered': agg['last_delivered'].strftime('%Y-%m-%d %H:%M') if agg['last_delivered'] else '',
            'by_stream': [{'stream': r['cdr_file__decoder_type'] or '-',
                           'files': r['files'], 'records': r['records'] or 0}
                          for r in by_stream],
        })

    # --- Summary ---
    total_in = file_qs.aggregate(files=Count('id'), records=Sum('records_valid'))
    total_out = dist_qs.filter(status=DistributionLog.Status.SUCCESS).aggregate(
        files=Count('id'), records=Sum('record_count'))
    completed = file_qs.filter(status=CDRFile.Status.COMPLETED).count()
    failed = file_qs.filter(status=CDRFile.Status.FAILED).count()
    success_rate = round(100 * completed / (completed + failed), 1) if (completed + failed) else 0

    # --- Stream-type breakdown (#2) ---
    stream_breakdown = list(
        file_qs.exclude(decoder_type='')
        .values('decoder_type')
        .annotate(files=Count('id'), records=Sum('records_valid'))
        .order_by('-files')
    )
    stream_data = [{'stream': r['decoder_type'] or 'UNKNOWN',
                    'files': r['files'], 'records': r['records'] or 0}
                   for r in stream_breakdown]

    # --- Error analysis (#3) ---
    error_file_ids = file_qs.values_list('id', flat=True)
    top_errors = list(
        ProcessingError.objects
        .filter(cdr_file_id__in=error_file_ids)
        .values('stage', 'error_class')
        .annotate(count=Count('id'))
        .order_by('-count')[:10]
    )
    file_errors = list(
        file_qs.filter(status=CDRFile.Status.FAILED)
        .exclude(error_message='')
        .values('error_message')
        .annotate(count=Count('id'))
        .order_by('-count')[:5]
    )
    error_analysis = {
        'processing_errors': [
            {'stage': e['stage'], 'error_class': e['error_class'] or '(unknown)',
             'count': e['count']}
            for e in top_errors
        ],
        'file_errors': [
            {'message': e['error_message'][:120], 'count': e['count']}
            for e in file_errors
        ],
        'total_processing_errors': ProcessingError.objects.filter(cdr_file_id__in=error_file_ids).count(),
        'total_failed_files': failed,
    }

    # --- Processing latency (#4) ---
    latency_qs = file_qs.filter(
        processing_started__isnull=False,
        processing_completed__isnull=False,
        status=CDRFile.Status.COMPLETED,
    ).annotate(
        proc_duration=ExpressionWrapper(
            F('processing_completed') - F('processing_started'),
            output_field=DurationField(),
        )
    )
    latency_agg = latency_qs.aggregate(
        avg_duration=Avg('proc_duration'),
        min_duration=Min('proc_duration'),
        max_duration=Max('proc_duration'),
        file_count=Count('id'),
    )
    def _dur_seconds(d):
        return round(d.total_seconds(), 1) if d else None
    latency = {
        'avg_seconds': _dur_seconds(latency_agg['avg_duration']),
        'min_seconds': _dur_seconds(latency_agg['min_duration']),
        'max_seconds': _dur_seconds(latency_agg['max_duration']),
        'sample_count': latency_agg['file_count'] or 0,
    }

    # --- Comparison with previous period (#5) ---
    prev_total_in = prev_file_qs.aggregate(files=Count('id'), records=Sum('records_valid'))
    prev_total_out = prev_dist_qs.filter(status=DistributionLog.Status.SUCCESS).aggregate(
        files=Count('id'), records=Sum('record_count'))
    prev_completed = prev_file_qs.filter(status=CDRFile.Status.COMPLETED).count()
    prev_failed = prev_file_qs.filter(status=CDRFile.Status.FAILED).count()
    prev_success_rate = round(100 * prev_completed / (prev_completed + prev_failed), 1) if (prev_completed + prev_failed) else 0

    def _delta(cur, prev):
        if prev and prev > 0:
            return round(((cur - prev) / prev) * 100, 1)
        return None

    comparison = {
        'prev_start': prev_start.strftime('%Y-%m-%d'),
        'prev_end': (prev_end - timedelta(days=1)).strftime('%Y-%m-%d'),
        'prev_files_in': prev_total_in['files'] or 0,
        'prev_records_in': prev_total_in['records'] or 0,
        'prev_files_out': prev_total_out['files'] or 0,
        'prev_records_out': prev_total_out['records'] or 0,
        'prev_success_rate': prev_success_rate,
        'delta_files_in': _delta(total_in['files'] or 0, prev_total_in['files'] or 0),
        'delta_records_in': _delta(total_in['records'] or 0, prev_total_in['records'] or 0),
        'delta_files_out': _delta(total_out['files'] or 0, prev_total_out['files'] or 0),
        'delta_records_out': _delta(total_out['records'] or 0, prev_total_out['records'] or 0),
        'delta_success_rate': round(success_rate - prev_success_rate, 1) if prev_success_rate is not None else None,
    }

    # --- Distribution rule performance (#8) ---
    rule_perf = []
    for rule in DistributionRule.objects.filter(is_active=True).order_by('priority', 'name'):
        r_logs = dist_qs.filter(rule=rule)
        r_agg = r_logs.aggregate(
            total=Count('id'),
            success=Count('id', filter=Q(status=DistributionLog.Status.SUCCESS)),
            failed=Count('id', filter=Q(status=DistributionLog.Status.FAILED)),
            skipped=Count('id', filter=Q(status=DistributionLog.Status.SKIPPED)),
            total_records=Sum('record_count'),
            total_retries=Sum('retry_count'),
        )
        if r_agg['total'] == 0:
            continue
        r_success_rate = round(100 * (r_agg['success'] or 0) / r_agg['total'], 1) if r_agg['total'] else 0
        rule_perf.append({
            'rule_id': rule.pk,
            'rule_name': rule.name,
            'stream_type': rule.stream_type,
            'portal_name': rule.output_portal.name if rule.output_portal else '-',
            'total': r_agg['total'] or 0,
            'success': r_agg['success'] or 0,
            'failed': r_agg['failed'] or 0,
            'skipped': r_agg['skipped'] or 0,
            'records': r_agg['total_records'] or 0,
            'retries': r_agg['total_retries'] or 0,
            'success_rate': r_success_rate,
        })

    # --- Daily chart ---
    by_day = list(
        file_qs.filter(status=CDRFile.Status.COMPLETED)
        .annotate(day=TruncDate('created_at'))
        .values('day')
        .annotate(files=Count('id'), records=Sum('records_valid'))
        .order_by('day')
    )

    return JsonResponse({
        'success': True,
        'summary': {
            'total_files_in': total_in['files'] or 0,
            'total_records_in': total_in['records'] or 0,
            'total_files_out': total_out['files'] or 0,
            'total_records_out': total_out['records'] or 0,
            'success_rate': success_rate,
        },
        'input_sources': input_sources,
        'output_portals': output_portals,
        'stream_breakdown': stream_data,
        'error_analysis': error_analysis,
        'latency': latency,
        'comparison': comparison,
        'rule_performance': rule_perf,
        'by_day': [{'day': r['day'].isoformat() if r['day'] else '',
                    'files': r['files'], 'records': r['records'] or 0}
                   for r in by_day],
    })


@login_required
def reports_export(request):
    if request.method != 'POST':
        return JsonResponse({'success': False, 'message': 'POST required'}, status=405)

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from django.db.models.functions import TruncDate
    from collection.models import DistributionLog

    start_date = request.POST.get('start_date', '').strip()
    end_date = request.POST.get('end_date', '').strip()

    if not start_date or not end_date:
        return JsonResponse({'success': False, 'message': 'Both dates required.'}, status=400)

    try:
        dt_start = datetime.strptime(start_date, '%Y-%m-%d')
        dt_end = datetime.strptime(end_date, '%Y-%m-%d') + timedelta(days=1)
    except ValueError:
        return JsonResponse({'success': False, 'message': 'Invalid date format.'}, status=400)

    operator = request.POST.get('operator', '').strip()

    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='0D6EFD', end_color='0D6EFD', fill_type='solid')
    green_fill = PatternFill(start_color='198754', end_color='198754', fill_type='solid')
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin'),
    )
    num_fmt = '#,##0'

    def style_header(ws, row, cols, fill=None):
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.fill = fill or header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

    def style_data(ws, row, cols):
        for col in range(1, cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border

    # --- Sheet 1: Summary ---
    ws = wb.active
    ws.title = 'Summary'
    file_qs = CDRFile.objects.filter(created_at__gte=dt_start, created_at__lt=dt_end)
    dist_qs = DistributionLog.objects.filter(delivered_at__gte=dt_start, delivered_at__lt=dt_end)
    if operator:
        file_qs = file_qs.filter(operator_code=operator)
        dist_qs = dist_qs.filter(cdr_file__operator_code=operator)

    total_in = file_qs.aggregate(files=Count('id'), records=Sum('records_valid'))
    total_out = dist_qs.filter(status=DistributionLog.Status.SUCCESS).aggregate(
        files=Count('id'), records=Sum('record_count'))
    completed = file_qs.filter(status=CDRFile.Status.COMPLETED).count()
    failed = file_qs.filter(status=CDRFile.Status.FAILED).count()
    success_rate = round(100 * completed / (completed + failed), 1) if (completed + failed) else 0

    ws.append(['UMP Mediation Report'])
    ws['A1'].font = Font(bold=True, size=14)
    ws.append([f'Period: {start_date} to {end_date}'])
    ws.append([])
    ws.append(['Metric', 'Value'])
    style_header(ws, 4, 2)
    for label, val in [
        ('Files Received', total_in['files'] or 0),
        ('Files Distributed', total_out['files'] or 0),
        ('Records Processed', total_in['records'] or 0),
        ('Records Delivered', total_out['records'] or 0),
        ('Success Rate', f'{success_rate}%'),
        ('Completed Files', completed),
        ('Failed Files', failed),
    ]:
        ws.append([label, val])
    for r in range(5, ws.max_row + 1):
        style_data(ws, r, 2)
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 18

    # --- Sheet 2: Input Sources ---
    ws2 = wb.create_sheet('Input Sources')
    ws2.append(['Source', 'Stream', 'Files', 'Records', 'Completed', 'Failed', 'Duplicate', 'First Received', 'Last Received'])
    style_header(ws2, 1, 9)

    sources = DataSource.objects.filter(enabled=True).order_by('name')
    row = 2
    for src in sources:
        src_files = file_qs.filter(source=src)
        agg = src_files.aggregate(
            total_files=Count('id'), total_records=Sum('records_valid'),
            completed=Count('id', filter=Q(status=CDRFile.Status.COMPLETED)),
            failed=Count('id', filter=Q(status=CDRFile.Status.FAILED)),
            duplicates=Count('id', filter=Q(status=CDRFile.Status.DUPLICATE)),
            first_received=Min('created_at'), last_received=Max('created_at'),
        )
        if agg['total_files'] == 0:
            continue
        by_stream = list(
            src_files.values('decoder_type')
            .annotate(files=Count('id'), records=Sum('records_valid'))
            .order_by('decoder_type')
        )
        first_dt = agg['first_received'].strftime('%Y-%m-%d %H:%M') if agg['first_received'] else ''
        last_dt = agg['last_received'].strftime('%Y-%m-%d %H:%M') if agg['last_received'] else ''
        if by_stream:
            for s in by_stream:
                ws2.append([src.name, s['decoder_type'] or '-', s['files'],
                            s['records'] or 0, '', '', '', '', ''])
                style_data(ws2, row, 9)
                row += 1
        ws2.append([f'{src.name} Total', '', agg['total_files'] or 0,
                     agg['total_records'] or 0, agg['completed'] or 0,
                     agg['failed'] or 0, agg['duplicates'] or 0, first_dt, last_dt])
        for col in range(1, 10):
            cell = ws2.cell(row=row, column=col)
            cell.font = Font(bold=True)
            cell.border = thin_border
        row += 1

    # No-source files
    no_src = file_qs.filter(source__isnull=True)
    no_agg = no_src.aggregate(
        total_files=Count('id'), total_records=Sum('records_valid'),
        completed=Count('id', filter=Q(status=CDRFile.Status.COMPLETED)),
        failed=Count('id', filter=Q(status=CDRFile.Status.FAILED)),
        duplicates=Count('id', filter=Q(status=CDRFile.Status.DUPLICATE)),
        first_received=Min('created_at'), last_received=Max('created_at'),
    )
    if no_agg['total_files']:
        no_src_streams = list(
            no_src.values('decoder_type')
            .annotate(files=Count('id'), records=Sum('records_valid'))
            .order_by('decoder_type')
        )
        first_dt = no_agg['first_received'].strftime('%Y-%m-%d %H:%M') if no_agg['first_received'] else ''
        last_dt = no_agg['last_received'].strftime('%Y-%m-%d %H:%M') if no_agg['last_received'] else ''
        if no_src_streams:
            for s in no_src_streams:
                ws2.append(['(No Source)', s['decoder_type'] or '-', s['files'],
                            s['records'] or 0, '', '', '', '', ''])
                style_data(ws2, row, 9)
                row += 1
        ws2.append(['(No Source) Total', '', no_agg['total_files'] or 0,
                     no_agg['total_records'] or 0, no_agg['completed'] or 0,
                     no_agg['failed'] or 0, no_agg['duplicates'] or 0, first_dt, last_dt])
        for col in range(1, 10):
            cell = ws2.cell(row=row, column=col)
            cell.font = Font(bold=True)
            cell.border = thin_border
        row += 1

    for col in range(1, 10):
        ws2.column_dimensions[get_column_letter(col)].width = 16
    ws2.column_dimensions['A'].width = 28
    ws2.column_dimensions['H'].width = 20
    ws2.column_dimensions['I'].width = 20

    # --- Sheet 3: Output Portals ---
    ws3 = wb.create_sheet('Output Portals')
    ws3.append(['Portal', 'Stream', 'Deliveries', 'Records', 'Success', 'Failed', 'Skipped', 'First Delivered', 'Last Delivered'])
    style_header(ws3, 1, 9, fill=green_fill)

    from portals.models import OutputPortal
    row = 2
    for portal in OutputPortal.objects.filter(is_active=True).order_by('name'):
        p_logs = dist_qs.filter(output_portal=portal)
        agg = p_logs.aggregate(
            total_files=Count('id'), total_records=Sum('record_count'),
            success=Count('id', filter=Q(status=DistributionLog.Status.SUCCESS)),
            failed=Count('id', filter=Q(status=DistributionLog.Status.FAILED)),
            skipped=Count('id', filter=Q(status=DistributionLog.Status.SKIPPED)),
            first_delivered=Min('delivered_at'), last_delivered=Max('delivered_at'),
        )
        if agg['total_files'] == 0:
            continue
        by_stream = list(
            p_logs.filter(cdr_file__isnull=False)
            .values('cdr_file__decoder_type')
            .annotate(files=Count('id'), records=Sum('record_count'))
            .order_by('cdr_file__decoder_type')
        )
        first_dt = agg['first_delivered'].strftime('%Y-%m-%d %H:%M') if agg['first_delivered'] else ''
        last_dt = agg['last_delivered'].strftime('%Y-%m-%d %H:%M') if agg['last_delivered'] else ''
        if by_stream:
            for s in by_stream:
                ws3.append([portal.name, s['cdr_file__decoder_type'] or '-',
                            s['files'], s['records'] or 0, '', '', '', '', ''])
                style_data(ws3, row, 9)
                row += 1
        ws3.append([f'{portal.name} Total', '', agg['total_files'] or 0,
                     agg['total_records'] or 0, agg['success'] or 0,
                     agg['failed'] or 0, agg['skipped'] or 0, first_dt, last_dt])
        for col in range(1, 10):
            cell = ws3.cell(row=row, column=col)
            cell.font = Font(bold=True)
            cell.border = thin_border
        row += 1

    for col in range(1, 10):
        ws3.column_dimensions[get_column_letter(col)].width = 16
    ws3.column_dimensions['A'].width = 28
    ws3.column_dimensions['H'].width = 20
    ws3.column_dimensions['I'].width = 20

    # --- Sheet 4: Daily Breakdown ---
    ws4 = wb.create_sheet('Daily Breakdown')
    ws4.append(['Date', 'Files', 'Records'])
    style_header(ws4, 1, 3)
    by_day = (
        file_qs.filter(status=CDRFile.Status.COMPLETED)
        .annotate(day=TruncDate('created_at'))
        .values('day')
        .annotate(files=Count('id'), records=Sum('records_valid'))
        .order_by('day')
    )
    row = 2
    for r in by_day:
        ws4.append([r['day'].isoformat() if r['day'] else '', r['files'], r['records'] or 0])
        style_data(ws4, row, 3)
        row += 1
    ws4.column_dimensions['A'].width = 14
    ws4.column_dimensions['B'].width = 12
    ws4.column_dimensions['C'].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    from django.http import HttpResponse
    fname = f'UMP_Report_{start_date}_to_{end_date}.xlsx'
    response = HttpResponse(buf.read(),
                            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{fname}"'
    return response


