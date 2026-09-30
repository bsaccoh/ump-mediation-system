"""Render assembled report data to an .xlsx workbook (openpyxl)."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill('solid', fgColor='F1F5F9')
_HEADER_FONT = Font(bold=True, color='334155')
_TITLE_FONT = Font(bold=True, size=14, color='0D47A1')


def _safe_sheet_name(name, used):
    clean = re.sub(r'[\\/*?:\[\]]', ' ', name)[:28].strip() or 'Sheet'
    candidate, i = clean, 1
    while candidate in used:
        i += 1
        candidate = f'{clean[:25]} {i}'
    used.add(candidate)
    return candidate


def _write_table(ws, columns, rows, start_row=1):
    for c, col in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=c, value=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for r, row in enumerate(rows, start=start_row + 1):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=('' if val is None else val))
    # Column widths from header + a sample of values.
    for c, col in enumerate(columns, start=1):
        width = max([len(str(col))] + [len(str(row[c - 1])) for row in rows[:50] if c - 1 < len(row)] + [8])
        ws.column_dimensions[get_column_letter(c)].width = min(width + 2, 40)
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def render_excel(data) -> bytes:
    wb = Workbook()
    used = set()

    ws = wb.active
    ws.title = _safe_sheet_name('Summary', used)
    ws['A1'] = data.get('title', 'Drive Test Report')
    ws['A1'].font = _TITLE_FONT
    meta = data.get('meta') or {}
    row = 3
    for k in ('campaign', 'project', 'operator', 'technology', 'region', 'generated_at', 'prepared_by'):
        if k in meta:
            ws.cell(row=row, column=1, value=k.replace('_', ' ').title()).font = _HEADER_FONT
            ws.cell(row=row, column=2, value=meta[k])
            row += 1
    row += 1
    summary = data.get('summary') or []
    if summary:
        rows = [[s['label'], s['value'], s.get('unit', '')] for s in summary]
        _write_table(ws, ['Metric', 'Value', 'Unit'], rows, start_row=row)

    for sec in data.get('sections') or []:
        if sec.get('kind') != 'table' or not sec.get('rows'):
            continue
        sheet = wb.create_sheet(_safe_sheet_name(sec['heading'], used))
        _write_table(sheet, sec['columns'], sec['rows'])
        if sec.get('histogram'):
            base = len(sec['rows']) + 3
            sheet.cell(row=base, column=1, value='Distribution').font = _HEADER_FONT
            _write_table(sheet, ['Class', 'Count', 'Share %'],
                         [[h['label'], h['count'], h['pct']] for h in sec['histogram']],
                         start_row=base + 1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
