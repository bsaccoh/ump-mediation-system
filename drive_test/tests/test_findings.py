"""Analysis / Findings page tests. Fixtures exist only in the isolated test DB."""
import io
from datetime import date, datetime

from django.contrib.auth import get_user_model
from django.test import Client, TestCase, override_settings
from django.urls import reverse

from drive_test.models import (
    Cell, DriveTestFile, DriveTestSession, Finding, Measurement, RadioMeasurement, RegulatoryRule,
    RegulatoryThreshold, Sector, ServiceMeasurement, Site,
)
from reference.models import Operator

STATIC = override_settings(STORAGES={
    'default': {'BACKEND': 'django.core.files.storage.FileSystemStorage'},
    'staticfiles': {'BACKEND': 'django.contrib.staticfiles.storage.StaticFilesStorage'}})


@STATIC
class FindingsTest(TestCase):
    databases = {'default'}

    @classmethod
    def setUpTestData(cls):
        U = get_user_model()
        cls.viewer = U.objects.create_user('viewer', password='x')
        cls.analyst = U.objects.create_user('analyst', password='x', is_analyst=True)
        cls.orange = Operator.objects.create(code='orange', name='Orange SL', home_plmn='61901', home_mcc='619', home_mnc='01')
        cls.qcell = Operator.objects.create(code='qcell', name='Qcell', home_plmn='61905', home_mcc='619', home_mnc='05')
        cls.site = Site.objects.create(site_id='FT-001', name='Freetown Central', operator=cls.orange, latitude=8.48, longitude=-13.23)
        cls.sec = Sector.objects.create(site=cls.site, sector_id='A')
        cls.cell = Cell.objects.create(cell_id='LTE-001', operator=cls.orange, sector=cls.sec, technology='4G')

        cls.sess = DriveTestSession.objects.create(operator=cls.orange, test_date='2026-09-24', uploaded_by=cls.viewer, status='COMPLETED')
        cls.sess2 = DriveTestSession.objects.create(operator=cls.qcell, test_date='2026-08-01', uploaded_by=cls.viewer, status='COMPLETED')
        cls.busy = DriveTestSession.objects.create(operator=cls.orange, test_date='2026-09-25', uploaded_by=cls.viewer, status='PROCESSING')
        f = DriveTestFile.objects.create(session=cls.sess, original_filename='walk_001.trp', file_path='p', file_size=1, sha256='a' * 64)
        f2 = DriveTestFile.objects.create(session=cls.sess2, original_filename='q.trp', file_path='p', file_size=1, sha256='b' * 64)

        def meas(file, seq, lat, lon, cell=None, tech='4G', **radio):
            m = Measurement.objects.create(drive_file=file, sequence_num=seq, captured_at=datetime(2026, 9, 24, 12, 0, seq),
                                           latitude=lat, longitude=lon, matched_cell=cell,
                                           match_method='exact_ecgi' if cell else '')
            RadioMeasurement.objects.create(measurement=m, technology=tech, **radio)
            return m

        cls.m1 = meas(f, 1, 8.4841, -13.2299, cls.cell, rsrp=-118.0)
        cls.m2 = meas(f, 2, 0.0, 0.0, None, tech='2G', rssi=-101.0)          # no valid location, unmatched
        cls.m3 = meas(f2, 3, 7.9, -11.7, None, tech='3G')
        ServiceMeasurement.objects.create(measurement=cls.m1, service_type='VOICE', outcome='DROPPED', call_duration_s=42)

        cls.rule = RegulatoryRule.objects.create(rule_code='LTE-RSRP', name='LTE RSRP Coverage', metric='rsrp', condition='lt',
                                                 effective_from=date(2025, 1, 1), authority='NATCA')
        cls.thr = RegulatoryThreshold.objects.create(rule=cls.rule, critical_value=-110.0, warning_value=-105.0, unit='dBm',
                                                     effective_from=date(2025, 1, 1))

        def fd(session, ftype, sev, desc, **kw):
            return Finding.objects.create(session=session, finding_type=ftype, severity=sev, description=desc, **kw)

        cls.f_high = fd(cls.sess, 'WEAK_SIGNAL', 'HIGH', 'RSRP below threshold', measurement=cls.m1, cell=cls.cell,
                        measured_value=-118.0, threshold_value=-110.0, threshold=cls.thr)
        cls.f_crit = fd(cls.sess, 'ANOMALY', 'CRITICAL', 'Missing GPS on 40% of points')
        cls.f_low = fd(cls.sess, 'DROP_CALL', 'LOW', 'Call dropped', measurement=cls.m2, is_resolved=True)
        cls.f_q = fd(cls.sess2, 'LOW_THROUGHPUT', 'MEDIUM', 'Slow download', measurement=cls.m3, latitude=7.9, longitude=-11.7)
        cls.url = reverse('drive_test:finding_list')

    def setUp(self):
        self.client.force_login(self.viewer)

    def ids(self, **q):
        return sorted(f.pk for f in self.client.get(self.url, q).context['findings'])

    def pks(self, *fs):
        return sorted(f.pk for f in fs)

    def test_summary_uses_backend_values(self):
        s = self.client.get(self.url).context['summary']
        self.assertEqual(s, {'total': 4, 'critical': 1, 'high': 1, 'open': 3})

    def test_filters(self):
        self.assertEqual(self.ids(operator='qcell'), self.pks(self.f_q))
        self.assertEqual(self.ids(technology='4G'), self.pks(self.f_high))          # via radio or cell
        self.assertEqual(self.ids(technology='2G'), self.pks(self.f_low))
        self.assertEqual(self.ids(severity='CRITICAL'), self.pks(self.f_crit))
        self.assertEqual(self.ids(category='DROP_CALL'), self.pks(self.f_low))
        self.assertEqual(self.ids(status='resolved'), self.pks(self.f_low))
        self.assertEqual(self.ids(status='open'), self.pks(self.f_high, self.f_crit, self.f_q))
        self.assertEqual(self.ids(session=self.sess2.session_ref[:8].lower()), self.pks(self.f_q))
        self.assertEqual(self.ids(date_from='2026-09-01'), self.pks(self.f_high, self.f_crit, self.f_low))
        self.assertEqual(self.ids(date_to='2026-08-31'), self.pks(self.f_q))
        self.assertEqual(self.ids(severity='bogus'), self.ids())                   # unknown values are ignored
        self.assertEqual(self.ids(date_from='garbage'), [])                        # never a 500

    def test_search_fields(self):
        self.assertEqual(self.ids(q='below threshold'), self.pks(self.f_high))
        self.assertEqual(self.ids(q='weak signal'), self.pks(self.f_high))         # category label
        self.assertEqual(self.ids(q='walk_001'), self.pks(self.f_high, self.f_low))  # source file
        self.assertEqual(self.ids(q='lte-001'), self.pks(self.f_high))             # cell
        self.assertEqual(self.ids(q='freetown'), self.pks(self.f_high))            # site
        self.assertEqual(self.ids(q=self.sess.session_ref), self.pks(self.f_high, self.f_crit, self.f_low))

    def test_row_data_is_not_reinterpreted(self):
        by = {f.pk: f for f in self.client.get(self.url).context['findings']}
        h = by[self.f_high.pk]
        self.assertEqual((h.severity, h.tech, h.eff_cell, h.point), ('HIGH', '4G', self.cell, (8.4841, -13.2299)))
        self.assertEqual(by[self.f_crit.pk].severity, 'CRITICAL')
        self.assertTrue(by[self.f_crit.pk].is_quality)                              # backend's own "Data Anomaly" type
        self.assertFalse(h.is_quality)
        self.assertEqual(by[self.f_low.pk].point, (None, None))                     # (0,0) is not a location
        self.assertIsNone(by[self.f_low.pk].eff_cell)
        self.assertEqual(by[self.f_q.pk].point, (7.9, -11.7))
        r = self.client.get(self.url)
        self.assertContains(r, 'Unmatched')
        self.assertNotContains(r, '0.0000, 0.0000')

    def test_matched_cell_used_when_finding_has_no_cell(self):
        Finding.objects.filter(pk=self.f_high.pk).update(cell=None)
        by = {f.pk: f for f in self.client.get(self.url).context['findings']}
        self.assertEqual(by[self.f_high.pk].eff_cell, self.cell)                    # real lineage via measurement

    def test_banners(self):
        r = self.client.get(self.url)
        self.assertContains(r, 'still processing')
        self.assertNotContains(r, 'No active regulatory rules')               # a rule is active in this fixture
        DriveTestSession.objects.filter(pk=self.busy.pk).update(status='COMPLETED')
        self.assertNotContains(self.client.get(self.url), 'still processing')

    def test_no_rules_banner_and_empty_state_is_honest(self):
        RegulatoryRule.objects.all().update(is_active=False)
        Finding.objects.all().delete()
        r = self.client.get(self.url)
        self.assertContains(r, 'No findings available')
        self.assertContains(r, 'does not mean the network passed')
        self.assertNotContains(r, 'Export')
        self.assertContains(self.client.get(self.url, {'q': 'zzz'}), 'No findings')

    def test_filtered_empty_and_pagination(self):
        self.assertContains(self.client.get(self.url, {'q': 'zzzz'}), 'No findings match your filters.')
        self.assertEqual(self.client.get(self.url, {'per_page': 7}).context['per_page'], 25)

    def test_no_n_plus_one(self):
        with self.assertNumQueries(12):
            self.client.get(self.url)

    def test_map_data(self):
        d = self.client.get(reverse('drive_test:finding_map_data')).json()
        self.assertEqual(sorted(f['properties']['id'] for f in d['features']), self.pks(self.f_high, self.f_q))
        self.assertEqual(self.client.get(reverse('drive_test:finding_map_data'), {'severity': 'LOW'}).json()['features'], [])

    def test_export_respects_filters(self):
        import openpyxl
        r = self.client.get(reverse('drive_test:finding_export'), {'severity': 'HIGH'})
        self.assertEqual(r.status_code, 200)
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        rows = list(ws.iter_rows(values_only=True))
        self.assertEqual(len(rows), 2)
        self.assertEqual((rows[1][2], rows[1][8], rows[1][15]), ('HIGH', 'LTE-001', 'LTE-RSRP'))
        self.assertEqual(len(list(openpyxl.load_workbook(io.BytesIO(self.client.get(reverse('drive_test:finding_export')).content)).active.iter_rows())), 5)

    def test_detail_evidence_rule_and_lineage(self):
        r = self.client.get(reverse('drive_test:finding_detail', args=[self.f_high.pk]))
        self.assertContains(r, 'LTE RSRP Coverage')
        self.assertContains(r, 'walk_001.trp')
        self.assertContains(r, reverse('drive_test:session_detail', args=[self.sess.session_ref]))
        self.assertContains(r, reverse('drive_test:sector_detail', args=[self.sec.pk]))
        self.assertContains(r, 'Dropped')                                            # service evidence
        self.assertEqual((r.context['lat'], r.context['technology'], r.context['rule_current']), (8.4841, '4G', True))
        self.assertIn(('RSRP', -118.0, 'dBm'), r.context['radio_values'])
        self.assertNotIn('RSRQ', [v[0] for v in r.context['radio_values']])        # absent values are not shown
        # missing threshold / measurement are stated, not invented
        r = self.client.get(reverse('drive_test:finding_detail', args=[self.f_crit.pk]))
        self.assertContains(r, 'No regulatory threshold configured')
        self.assertContains(r, 'No measurement or measured value')
        self.assertContains(r, 'No location is recorded')

    def test_expired_threshold_not_shown_as_current(self):
        RegulatoryThreshold.objects.filter(pk=self.thr.pk).update(effective_to=date(2025, 12, 31))
        r = self.client.get(reverse('drive_test:finding_detail', args=[self.f_high.pk]))
        self.assertFalse(r.context['rule_current'])
        self.assertContains(r, 'no longer effective')

    def test_status_workflow_permissions_and_immutability(self):
        url = reverse('drive_test:finding_set_status', args=[self.f_high.pk])
        self.assertEqual(self.client.get(url).status_code, 405)                     # POST only
        self.client.post(url, {'action': 'resolve'})                                # viewer: no effect
        self.f_high.refresh_from_db(); self.assertFalse(self.f_high.is_resolved)

        self.client.force_login(self.analyst)
        self.client.post(url, {'action': 'resolve', 'note': 'Fixed by retune', 'severity': 'INFO',
                               'description': 'tampered', 'measured_value': '0'})
        self.f_high.refresh_from_db()
        self.assertEqual((self.f_high.is_resolved, self.f_high.resolved_by, self.f_high.severity), (True, self.analyst, 'HIGH'))
        self.assertEqual((self.f_high.description, self.f_high.measured_value), ('RSRP below threshold', -118.0))
        self.assertIn('Fixed by retune', self.f_high.notes)
        self.client.post(url, {'action': 'resolve'})                                # already resolved: rejected
        self.assertEqual(self.f_high.notes.count('Resolved'), 1)
        self.client.post(url, {'action': 'reopen'})
        self.f_high.refresh_from_db()
        self.assertEqual((self.f_high.is_resolved, self.f_high.resolved_at, self.f_high.resolved_by), (False, None, None))
        self.assertEqual(self.f_high.notes.count('\n'), 1)                           # history kept, appended
        self.client.post(url, {'action': 'explode'})
        self.assertFalse(Finding.objects.get(pk=self.f_high.pk).is_resolved)

    def test_csrf_enforced_on_status_change(self):
        c = Client(enforce_csrf_checks=True); c.force_login(self.analyst)
        self.assertEqual(c.post(reverse('drive_test:finding_set_status', args=[self.f_high.pk]), {'action': 'resolve'}).status_code, 403)

    def test_login_and_404(self):
        self.assertEqual(self.client.get(reverse('drive_test:finding_detail', args=[999999])).status_code, 404)
        self.client.logout()
        for n, a in (('finding_list', []), ('finding_map_data', []), ('finding_export', []), ('finding_detail', [self.f_high.pk])):
            self.assertEqual(self.client.get(reverse(f'drive_test:{n}', args=a)).status_code, 302)
